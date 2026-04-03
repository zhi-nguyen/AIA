import asyncio
import websockets
import json
import logging
import tkinter as tk
from tkinter import messagebox
import tkinter.simpledialog as simpledialog
import os
import subprocess
import platform
import sys
import threading
import pystray
from PIL import Image, ImageDraw

# Optional dependencies for document generation
try:
    import docx
except ImportError:
    docx = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

def load_config():
    """Load config.json from the same directory as the executable/script."""
    if getattr(sys, 'frozen', False):
        # If the application is run as a bundle (PyInstaller)
        base_dir = os.path.dirname(sys.executable)
    else:
        # If the application is run from a Python script
        base_dir = os.path.dirname(os.path.abspath(__file__))
    
    config_path = os.path.join(base_dir, "config.json")
    try:
        with open(config_path, "r", encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError:
        logging.error(f"Cannot find {config_path}. Please place config.json next to the executable.")
        root = tk.Tk()
        root.withdraw()
        messagebox.showerror("Error", "config.json not found! Please ensure it is in the same folder as AIA_Agent.exe")
        root.destroy()
        sys.exit(1)
    except Exception as e:
        logging.error(f"Error reading config.json: {e}")
        root = tk.Tk()
        root.withdraw()
        messagebox.showerror("Error", f"Error reading config.json: {e}")
        root.destroy()
        sys.exit(1)

CONFIG = load_config()
USER_ID = CONFIG.get("user_id")
TOKEN = CONFIG.get("user_token")
SERVER_URL = CONFIG.get("server_url", "ws://localhost:8000/api/v1/ws/agent")

if not USER_ID or not TOKEN:
    logging.error("Missing user_id or user_token in config.json")
    root = tk.Tk()
    root.withdraw()
    messagebox.showerror("Error", "Missing user_id or user_token in config.json")
    root.destroy()
    sys.exit(1)

WS_URL = f"{SERVER_URL}/{USER_ID}?token={TOKEN}"

def open_file(file_path: str):
    """Automatically open the file after creation."""
    try:
        if platform.system() == 'Windows':
            os.startfile(file_path)
        elif platform.system() == 'Darwin':  # macOS
            subprocess.Popen(['open', file_path])
        else:  # Linux
            subprocess.Popen(['xdg-open', file_path])
        logging.info(f"Opened file: {file_path}")
    except Exception as e:
        logging.error(f"Failed to open file automatically: {e}")

def get_desktop_path() -> str:
    """Helper to get the user's Desktop path."""
    if platform.system() == "Windows":
        return os.path.join(os.environ.get("USERPROFILE", os.path.expanduser("~")), "Desktop")
    else:
        return os.path.join(os.path.expanduser("~"), "Desktop")

def ask_user_consent(action_desc: str) -> bool:
    """Show a Tkinter popup to ask for user consent."""
    root = tk.Tk()
    root.withdraw()
    root.attributes('-topmost', True)
    
    msg = f"AIA Assistant wants to {action_desc}. Do you allow this?"
    result = messagebox.askyesno("Security Prompt - AIA Assistant", msg)
    
    root.destroy()
    return result

async def handle_word_automation(data: dict) -> dict:
    if not docx:
        return {"status": "error", "message": "python-docx is not installed. Please run `pip install python-docx`."}
        
    if not ask_user_consent("create and populate data into the Report file (Word)"):
        return {"status": "rejected", "message": "User denied consent for word document creation."}
        
    inner_payload = data.get("data", {})
    template_name = inner_payload.get("template_name", "Report")
    doc_data = inner_payload.get("data", {})
    
    # Generate filename (e.g., Contract_NguyenVanA.docx)
    name_suffix = doc_data.get("Name", "")
    filename_base = f"{template_name}_{name_suffix}" if name_suffix else f"AIA_{template_name}"
    filename_base = filename_base.replace(" ", "")
    file_path = os.path.join(get_desktop_path(), f"{filename_base}.docx")
    
    try:
        template_file = f"{template_name}.docx"
        if os.path.exists(template_file):
            doc = docx.Document(template_file)
        else:
            doc = docx.Document()
            # If no actual template file is found on disk, dynamically generate a formatted document
            for key, value in doc_data.items():
                key_lower = str(key).lower()
                clean_key = str(key).replace("_", " ").title()
                
                if key_lower == "title" or key_lower == "tieu de" or key_lower == "tiêu đề":
                    doc.add_heading(str(value), level=0)
                elif key_lower == "subtitle" or key_lower == "phu de" or key_lower == "phụ đề":
                    doc.add_heading(str(value), level=1)
                elif "title" in key_lower or "tieu de" in key_lower or "tiêu đề" in key_lower:
                    doc.add_heading(str(value), level=2)
                else:
                    # Thoroughly hide generic structural JSON keys from being printed as Headings
                    hide_heading = False
                    vn_en_structural_keys = ["content", "body", "description", "conclusion", "author", "section", "paragraph", "data", "noi dung", "nội dung", "mo dau", "mở đầu", "ket luan", "kết luận", "nguoi", "người", "phan", "phần", "muc", "mục", "tac gia", "tác giả"]
                    for hide_word in vn_en_structural_keys:
                        if hide_word in key_lower:
                            hide_heading = True
                            break
                            
                    if not hide_heading:
                        doc.add_heading(clean_key, level=3)
                        
                    for line in str(value).split('\n'):
                        if line.strip():
                            if line.strip().startswith("-") or line.strip().startswith("*"):
                                clean_line = line.strip()[1:].strip()
                                doc.add_paragraph(clean_line, style='List Bullet')
                            else:
                                doc.add_paragraph(line.strip())
        
        # Iterate through paragraphs, locate placeholders like {{Name}}, and replace
        for p in doc.paragraphs:
            for key, value in doc_data.items():
                placeholder = f"{{{{{key}}}}}"
                if placeholder in p.text:
                    p.text = p.text.replace(placeholder, "")
                    text_to_fill = str(value)
                    
                    if '\n' in text_to_fill:
                        lines = text_to_fill.split('\n')
                        for i, line in enumerate(lines):
                            run = p.add_run(line.strip())
                            if i < len(lines) - 1:
                                run.add_break()
                    else:
                        p.add_run(text_to_fill)
        
        doc.save(file_path)
        open_file(file_path)
        
        return {"status": "success", "message": f"Successfully created {file_path}"}
    except Exception as e:
        return {"status": "error", "message": str(e)}

async def handle_excel_automation(data: dict) -> dict:
    if not openpyxl:
        return {"status": "error", "message": "openpyxl is not installed. Please run `pip install openpyxl`."}
        
    if not ask_user_consent("create and populate data into the Report file (Excel)"):
        return {"status": "rejected", "message": "User denied consent for excel document creation."}
        
    file_path = os.path.join(get_desktop_path(), "Report.xlsx")
    
    try:
        raw_payload = data.get("data", [])
        content_data = []
        headers = []
        
        if isinstance(raw_payload, dict):
            content_data = raw_payload.get("data", [])
            headers = raw_payload.get("headers", [])
        elif isinstance(raw_payload, list):
            content_data = raw_payload
            
        wb = openpyxl.Workbook()
        ws = wb.active
        
        # Determine headers
        if not headers and isinstance(content_data, list) and len(content_data) > 0 and isinstance(content_data[0], dict):
            headers = list(content_data[0].keys())
        if not headers and isinstance(content_data, list) and len(content_data) > 0 and isinstance(content_data[0], dict):
            headers = list(content_data[0].keys())

        # Write headers
        if headers:
            for col_idx, header in enumerate(headers, start=1):
                ws.cell(row=1, column=col_idx, value=header)
        else:
            # Fallback based on instructions if completely empty
            ws['A1'] = "Name"
            ws['B1'] = "Description"
            
        # Use a loop to populate subsequent data rows
        if isinstance(content_data, list):
            start_row = 2 if headers or ws['A1'].value else 1
            for i, row in enumerate(content_data, start=start_row):
                if isinstance(row, dict):
                    if headers:
                        for col_idx, key in enumerate(headers, start=1):
                            ws.cell(row=i, column=col_idx, value=row.get(key, ""))
                    else:
                        ws.cell(row=i, column=1, value=str(row))
                elif isinstance(row, list):
                    for j, val in enumerate(row, start=1):
                        ws.cell(row=i, column=j, value=val)
                else:
                    ws.cell(row=i, column=1, value=row)

        wb.save(file_path)
        open_file(file_path)
        
        return {"status": "success", "message": f"Successfully created {file_path}"}
    except Exception as e:
        return {"status": "error", "message": str(e)}

async def execute_action(data: dict) -> dict:
    action = data.get("action")
    if action == "create_word":
        return await handle_word_automation(data)
    elif action == "create_excel":
        return await handle_excel_automation(data)
    else:
        return {"status": "error", "message": f"Unknown action: {action}"}

def prompt_for_new_token():
    root = tk.Tk()
    root.title("AIA Agent - Cập nhật Token")
    
    window_width = 400
    window_height = 150
    screen_width = root.winfo_screenwidth()
    screen_height = root.winfo_screenheight()
    x = int((screen_width / 2) - (window_width / 2))
    y = int((screen_height / 2) - (window_height / 2))
    root.geometry(f"{window_width}x{window_height}+{x}+{y}")
    root.attributes('-topmost', True)
    
    new_token_var = tk.StringVar()
    
    tk.Label(root, text="Kết nối thất bại 5 lần (Token có thể đã hết hạn).\nVui lòng nhập token mới từ website:").pack(pady=10)
    entry = tk.Entry(root, textvariable=new_token_var, width=50)
    entry.pack(pady=5)
    
    def on_save():
        root.quit()
        
    def on_close():
        new_token_var.set("")
        root.quit()
    
    btn_frame = tk.Frame(root)
    btn_frame.pack(pady=10)
    
    tk.Button(btn_frame, text="Lưu", command=on_save, bg="#4CAF50", fg="white", width=10).pack(side=tk.LEFT, padx=10)
    tk.Button(btn_frame, text="Hủy", command=on_close, width=10).pack(side=tk.LEFT)
    
    root.protocol("WM_DELETE_WINDOW", on_close)
    root.mainloop()
    
    val = new_token_var.get().strip()
    root.destroy()
    return val if val else None

def update_config(new_token):
    global TOKEN, WS_URL
    if not new_token:
        return
    TOKEN = new_token
    base_dir = os.path.dirname(sys.executable) if getattr(sys, 'frozen', False) else os.path.dirname(os.path.abspath(__file__))
    config_path = os.path.join(base_dir, "config.json")
    try:
        with open(config_path, "r", encoding="utf-8") as f:
            config = json.load(f)
        config["user_token"] = new_token
        with open(config_path, "w", encoding="utf-8") as f:
            json.dump(config, f, indent=4)
        WS_URL = f"{SERVER_URL}/{USER_ID}?token={TOKEN}"
        logging.info("Token updated successfully. Restarting application...")
        
        args = [arg for arg in sys.argv[1:] if arg != '--prompt-token']
        # Restart the process
        if getattr(sys, 'frozen', False):
            subprocess.Popen([sys.executable] + args)
        else:
            subprocess.Popen([sys.executable, sys.argv[0]] + args)
        os._exit(0)
    except Exception as e:
        logging.error(f"Failed to update config.json or restart: {e}")

async def agent_loop() -> None:
    retry_count = 0
    while True:
        try:
            logging.info(f"Connecting to {WS_URL}...")
            async with websockets.connect(WS_URL) as ws:
                logging.info("Connected to AIA Backend!")
                retry_count = 0
                while True:
                    message: str = await ws.recv()
                    data: dict = json.loads(message)
                    logging.info(f"Received instruction: {data}")
                    
                    # Dispatch command
                    result = await execute_action(data)
                    
                    # Send response back to backend
                    response: dict = {
                        "action_id": data.get("action_id"),
                        "action": data.get("action"),
                        "result": result
                    }
                    await ws.send(json.dumps(response))
                    logging.info(f"Sent response: {response}")
        
        except websockets.exceptions.ConnectionClosedError as e:
            logging.error(f"WebSocket connection closed unexpectedly: {e}")
        except websockets.exceptions.InvalidStatusCode as e:
            logging.error(f"WebSocket auth failed: {e}")
        except ConnectionRefusedError:
            logging.error("Connection refused. Is the backend running?")
        except Exception as e:
            logging.error(f"Unexpected error: {e}")
            
        retry_count += 1
        if retry_count >= 5:
            logging.info("Connection failed 5 times. Launching token prompt...")
            if getattr(sys, 'frozen', False):
                subprocess.Popen([sys.executable, '--prompt-token'])
            else:
                subprocess.Popen([sys.executable, sys.argv[0], '--prompt-token'])
            os._exit(0)

        logging.info("Retrying in 5 seconds...")
        await asyncio.sleep(5)

def create_default_icon():
    """Generates a temporary placeholder icon (Blue square with yellow text)."""
    image = Image.new('RGB', (64, 64), color=(73, 109, 137))
    d = ImageDraw.Draw(image)
    d.text((15, 25), "AIA", fill=(255, 255, 0))
    return image

def on_quit(icon, item):
    """Callback to terminate the program when 'Quit' is selected."""
    icon.stop()  # Stop the System Tray loop
    os._exit(0)  # Immediate termination of the entire process (including background agent_loop)

def setup_tray():
    """Configures the context menu and initializes the System Tray."""
    # Define the right-click context menu
    menu = pystray.Menu(
        pystray.MenuItem("Quit AIA Assistant", on_quit)
    )
    
    # Initialize the icon (Identifier, Image, Hover text, Menu)
    icon = pystray.Icon("AIA_Agent", create_default_icon(), "AIA Assistant", menu)
    
    # Start the UI loop (This blocks the main thread until icon.stop() is called)
    icon.run()

if __name__ == "__main__":
    if '--prompt-token' in sys.argv:
        new_token = prompt_for_new_token()
        if new_token:
            update_config(new_token)
        sys.exit(0)

    # Offload the AI Agent event loop to a background thread.
    # daemon=True ensures the background thread exits when the main thread (Tray) terminates.
    agent_thread = threading.Thread(target=lambda: asyncio.run(agent_loop()), daemon=True)
    agent_thread.start()

    # Launch the System Tray on the Main Thread
    setup_tray()

