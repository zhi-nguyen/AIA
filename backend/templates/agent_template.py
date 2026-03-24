import asyncio
import websockets
import json
import logging
import tkinter as tk
from tkinter import messagebox
import tkinter.simpledialog as simpledialog
import os
import subprocess
import platform

# Optional dependencies for document generation
try:
    import docx
except ImportError:
    docx = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

# Credentials replaced by the Backend Download API
USER_ID = "{{USER_ID}}"
TOKEN = "{{USER_TOKEN}}"
WS_URL = f"ws://localhost:8000/api/v1/ws/agent/{USER_ID}?token={TOKEN}"


def open_file(file_path: str):
    """Automatically open the file after creation."""
    try:
        if platform.system() == 'Windows':
            os.startfile(file_path)
        elif platform.system() == 'Darwin':  # macOS
            subprocess.Popen(['open', file_path])
        else:  # Linux
            subprocess.Popen(['xdg-open', file_path])
        logging.info(f"Opened file: {file_path}")
    except Exception as e:
        logging.error(f"Failed to open file automatically: {e}")


def get_desktop_path() -> str:
    """Helper to get the user's Desktop path."""
    if platform.system() == "Windows":
        return os.path.join(os.environ.get("USERPROFILE", os.path.expanduser("~")), "Desktop")
    else:
        return os.path.join(os.path.expanduser("~"), "Desktop")


def ask_user_consent(action_desc: str) -> bool:
    """Show a Tkinter popup to ask for user consent."""
    root = tk.Tk()
    root.withdraw()
    root.attributes('-topmost', True)
    
    msg = f"AIA Assistant wants to {action_desc}. Do you allow this?"
    result = messagebox.askyesno("Security Prompt - AIA Assistant", msg)
    
    root.destroy()
    return result


async def handle_word_automation(data: dict) -> dict:
    if not docx:
        return {"status": "error", "message": "python-docx is not installed. Please run `pip install python-docx`."}
        
    if not ask_user_consent("create and populate data into the Report file (Word)"):
        return {"status": "rejected", "message": "User denied consent for word document creation."}
        
    template_name = data.get("template_name", "Report")
    doc_data = data.get("data", {})
    
    # Generate filename (e.g., Contract_NguyenVanA.docx)
    name_suffix = doc_data.get("Name", "")
    filename_base = f"{template_name}_{name_suffix}" if name_suffix else f"AIA_{template_name}"
    filename_base = filename_base.replace(" ", "")
    file_path = os.path.join(get_desktop_path(), f"{filename_base}.docx")
    
    try:
        template_file = f"{template_name}.docx"
        if os.path.exists(template_file):
            doc = docx.Document(template_file)
        else:
            doc = docx.Document()
            # If no actual template file is found on disk, we just mock replacing by appending
            for key, value in doc_data.items():
                doc.add_paragraph(f"{key}: {value}")
        
        # Iterate through paragraphs, locate placeholders like {{Name}}, and replace
        for paragraph in doc.paragraphs:
            for key, value in doc_data.items():
                placeholder = f"{{{{{key}}}}}"
                if placeholder in paragraph.text:
                    paragraph.text = paragraph.text.replace(placeholder, str(value))
        
        doc.save(file_path)
        open_file(file_path)
        
        return {"status": "success", "message": f"Successfully created {file_path}"}
    except Exception as e:
        return {"status": "error", "message": str(e)}


async def handle_excel_automation(data: dict) -> dict:
    if not openpyxl:
        return {"status": "error", "message": "openpyxl is not installed. Please run `pip install openpyxl`."}
        
    if not ask_user_consent("create and populate data into the Report file (Excel)"):
        return {"status": "rejected", "message": "User denied consent for excel document creation."}
        
    file_path = os.path.join(get_desktop_path(), "Report.xlsx")
    content_data = data.get("data", [])
    
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        
        # Assign values to cells if specified in data payload
        if "headers" in data:
            for col_idx, header in enumerate(data["headers"], start=1):
                ws.cell(row=1, column=col_idx, value=header)
        else:
            # Fallback based on instructions
            ws['A1'] = "Name"
            ws['B1'] = "Revenue"
            
        # Use a loop to populate subsequent data rows
        if isinstance(content_data, list):
            start_row = 2 if "headers" in data or ws['A1'].value else 1
            for i, row in enumerate(content_data, start=start_row):
                if isinstance(row, dict):
                    # We assume specific keys or just put them in cells
                    ws.cell(row=i, column=1, value=row.get("Name", ""))
                    ws.cell(row=i, column=2, value=row.get("Revenue", ""))
                elif isinstance(row, list):
                    for j, val in enumerate(row, start=1):
                        ws.cell(row=i, column=j, value=val)
                else:
                    ws.cell(row=i, column=1, value=row)

        wb.save(file_path)
        open_file(file_path)
        
        return {"status": "success", "message": f"Successfully created {file_path}"}
    except Exception as e:
        return {"status": "error", "message": str(e)}


async def execute_action(data: dict) -> dict:
    action = data.get("action")
    if action == "create_word":
        return await handle_word_automation(data)
    elif action == "create_excel":
        return await handle_excel_automation(data)
    else:
        return {"status": "error", "message": f"Unknown action: {action}"}


def prompt_for_new_token():
    root = tk.Tk()
    root.withdraw()
    root.attributes('-topmost', True)
    messagebox.showwarning("Connection Failed", "Connection failed 5 times. The token may have expired or was modified. Please generate a new token from the web interface and enter it here.", parent=root)
    new_token = simpledialog.askstring("Input", "Enter new AIA Agent token:", parent=root)
    root.destroy()
    return new_token

def update_globals(new_token):
    global TOKEN, WS_URL
    if not new_token:
        return
    TOKEN = new_token
    # In template mode the host might be hardcoded
    WS_URL = f"ws://localhost:8000/api/v1/ws/agent/{USER_ID}?token={TOKEN}"
    logging.info("Token updated successfully.")

async def agent_loop() -> None:
    retry_count = 0
    while True:
        try:
            logging.info(f"Connecting to {WS_URL}...")
            async with websockets.connect(WS_URL) as ws:
                logging.info("Connected to AIA Backend!")
                retry_count = 0
                while True:
                    message: str = await ws.recv()
                    data: dict = json.loads(message)
                    logging.info(f"Received instruction: {data}")
                    
                    # Dispatch command
                    result = await execute_action(data)
                    
                    # Send response back to backend
                    response: dict = {
                        "action_id": data.get("action_id"),
                        "action": data.get("action"),
                        "result": result
                    }
                    await ws.send(json.dumps(response))
                    logging.info(f"Sent response: {response}")
        
        except websockets.exceptions.ConnectionClosedError as e:
            logging.error(f"WebSocket connection closed unexpectedly: {e}")
        except websockets.exceptions.InvalidStatusCode as e:
            logging.error(f"WebSocket auth failed: {e}")
        except ConnectionRefusedError:
            logging.error("Connection refused. Is the backend running?")
        except Exception as e:
            logging.error(f"Unexpected error: {e}")
            
        retry_count += 1
        if retry_count >= 5:
            new_token = prompt_for_new_token()
            if new_token:
                update_globals(new_token)
            retry_count = 0
            
        logging.info("Retrying in 5 seconds...")
        await asyncio.sleep(5)


if __name__ == "__main__":
    asyncio.run(agent_loop())
